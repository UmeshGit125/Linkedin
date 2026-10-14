"""The Jobs reference sheet (Excel).

Tabs: "Top matches" (fit score >= jobs.min_score — the shortlist to apply from) and "Jobs"
(everything found). New jobs are appended at the bottom; rows already in the sheet are never
rewritten, so whatever you type in "My status" / "My notes" is kept."""
import os

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import CFG, ROOT

# (header, db column or callable, width)
COLUMNS = [
    ("Found at", lambda r: r["found_at"][:16].replace("T", " "), 16),
    ("Posted", lambda r: r["posted_ago"] or (r["posted_at"] or "")[:10], 14),
    ("Role", "title", 34),
    ("Company", "company", 22),
    ("Location", "location", 24),
    ("Work mode", "work_mode", 10),
    ("Employment type", "employment_type", 14),
    ("Seniority", "seniority", 14),
    ("Experience req.", "experience", 12),
    ("Tech stack", "tech_stack", 40),
    ("Salary", "salary", 18),
    ("Applicants", "applicants", 16),
    ("Industry", "industries", 22),
    ("Fit score", "score", 8),
    ("Fit notes", "reasons", 50),
    ("Keyword match", lambda r: "No" if r["status"] == "filtered" else "Yes", 9),
    ("Apply link", "url", 16),
    ("External apply", "apply_url", 16),
    ("Company page", "company_url", 16),
    ("Source", "source", 9),
    ("Job ID", "id", 7),
    ("My status", lambda r: "", 14),
    ("My notes", lambda r: "", 30),
]
LINK_COLUMNS = {"Apply link", "External apply", "Company page"}
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SCORE_FILLS = [(75, "C6EFCE"), (60, "FFEB9C"), (0, "F2F2F2")]


def path():
    return ROOT / CFG["jobs"].get("sheet_path", "Jobs.xlsx")


def _excel_has_it_open() -> bool:
    p = path()
    return (p.parent / f"~${p.name}").exists()


MATCHES = "Top matches"


def _setup_tab(ws) -> None:
    ws.append([c[0] for c in COLUMNS])
    for i, (_, _, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "D2"  # header + Found/Posted/Role stay visible while scrolling
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"


def _new_workbook() -> Workbook:
    wb = Workbook()
    wb.active.title = "Jobs"
    _setup_tab(wb.active)
    return wb


def _append_row(ws, r) -> None:
    values = [(col(r) if callable(col) else r[col]) for _, col, _ in COLUMNS]
    ws.append(["" if v is None else v for v in values])
    row_idx = ws.max_row
    for i, (header, _, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=row_idx, column=i)
        cell.alignment = Alignment(vertical="top", wrap_text=header in ("Fit notes", "Tech stack", "Role"))
        if header in LINK_COLUMNS and cell.value:
            cell.hyperlink = str(cell.value)
            cell.value = "Apply" if header == "Apply link" else "Open"
            cell.font = Font(color="0563C1", underline="single")
        if header == "Fit score":
            _paint_score(cell)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{ws.max_row}"


def _paint_score(cell) -> None:
    if isinstance(cell.value, int):
        color = next(c for threshold, c in SCORE_FILLS if cell.value >= threshold)
        cell.fill = PatternFill("solid", fgColor=color)


def _fill_scores(conn, ws) -> list[int]:
    """Writes Fit score / Fit notes into rows that were added before scoring finished.
    Only those two cells are touched — your own columns are never rewritten."""
    pending = {
        r["id"]: r for r in conn.execute(
            "SELECT id, score, reasons FROM jobs WHERE in_sheet = 1 AND score_in_sheet = 0 AND score IS NOT NULL"
        )
    }
    if not pending:
        return []
    headers = [c.value for c in ws[1]]
    id_col, score_col, notes_col = (headers.index(h) + 1 for h in ("Job ID", "Fit score", "Fit notes"))
    done = []
    for row in range(2, ws.max_row + 1):
        r = pending.get(ws.cell(row=row, column=id_col).value)
        if r is None:
            continue
        score_cell = ws.cell(row=row, column=score_col)
        score_cell.value = r["score"]
        _paint_score(score_cell)
        notes = ws.cell(row=row, column=notes_col)
        notes.value = r["reasons"]
        notes.alignment = Alignment(vertical="top", wrap_text=True)
        done.append(r["id"])
    return done


def sync(conn) -> int:
    """Appends jobs not yet in the sheet and fills in newly computed fit scores.
    Returns how many rows were added."""
    p = path()
    if _excel_has_it_open():
        print(f"[sheet] {p.name} is open in Excel — will update it on the next run.")
        return 0
    if p.exists():
        wb = load_workbook(p)
    else:
        wb = _new_workbook()
        conn.execute("UPDATE jobs SET in_sheet = 0, score_in_sheet = 0, match_in_sheet = 0")  # new sheet: full history
    ws = wb["Jobs"]
    if MATCHES not in wb.sheetnames:
        _setup_tab(wb.create_sheet(MATCHES, 0))  # first tab, so it opens on the shortlist
        conn.execute("UPDATE jobs SET match_in_sheet = 0")
    top = wb[MATCHES]

    scored_ids = _fill_scores(conn, ws)
    rows = conn.execute("SELECT * FROM jobs WHERE in_sheet = 0 ORDER BY id").fetchall()
    matches = conn.execute(
        "SELECT * FROM jobs WHERE match_in_sheet = 0 AND score >= ? ORDER BY score DESC, id",
        (CFG["jobs"]["min_score"],),
    ).fetchall()
    if not rows and not scored_ids and not matches:
        return 0
    for r in rows:
        _append_row(ws, r)
    for r in matches:
        _append_row(top, r)
    wb.active = 0

    tmp = p.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    os.replace(tmp, p)  # atomic: a crash mid-save never corrupts the sheet
    conn.executemany(
        "UPDATE jobs SET in_sheet = 1, score_in_sheet = (score IS NOT NULL) WHERE id = ?", [(r["id"],) for r in rows]
    )
    conn.executemany("UPDATE jobs SET score_in_sheet = 1 WHERE id = ?", [(i,) for i in scored_ids])
    conn.executemany("UPDATE jobs SET match_in_sheet = 1 WHERE id = ?", [(r["id"],) for r in matches])
    conn.commit()
    return len(rows)
